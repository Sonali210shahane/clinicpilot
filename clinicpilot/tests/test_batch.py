from pathlib import Path

from openpyxl import Workbook, load_workbook

from batch_triage import main, normalise

SAMPLE = Path(__file__).resolve().parent.parent / "sample_data" / "appointment_requests.csv"


def test_batch_report(tmp_path):
    output = tmp_path / "report.xlsx"
    assert main([str(SAMPLE), "-o", str(output)]) == 0

    wb = load_workbook(output)
    assert wb.sheetnames == ["Triage queue", "Summary", "Needs review"]

    rows = list(wb["Triage queue"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 23
    priorities = [r[1] for r in rows]
    order = ["Emergency", "Same-day", "Routine", "Telehealth", "Self-care"]
    assert priorities == sorted(priorities, key=order.index)
    assert priorities.count("Emergency") == 4

    review = list(wb["Needs review"].iter_rows(min_row=2, values_only=True))
    assert review[0][1] == "Unknown caller"


def test_batch_reads_excel_input(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "age", "request_type", "symptoms", "duration", "severity"])
    ws.append(["Sam", 29, "New problem", "Runny nose", "1-3 days", 2])
    ws.append(["Joan", 68, "Repeat script", "BP tablets", "n/a", 0])
    source = tmp_path / "requests.xlsx"
    wb.save(source)

    output = tmp_path / "out.xlsx"
    assert main([str(source), "-o", str(output)]) == 0
    rows = list(load_workbook(output)["Triage queue"].iter_rows(min_row=2, values_only=True))
    assert [r[1] for r in rows] == ["Telehealth", "Self-care"]


def test_missing_file():
    assert main(["does_not_exist.csv"]) == 1


def test_normalise_accepts_friendly_labels():
    row = normalise({"name": "A", "age": "40", "request_type": "Repeat prescription",
                     "duration": "More than a week", "pregnant": "Yes", "chronic": "no"})
    assert row["request_type"] == "script"
    assert row["duration"] == "gt7"
    assert row["pregnant"] is True
    assert row["chronic"] is False
