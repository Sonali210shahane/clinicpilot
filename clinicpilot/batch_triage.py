"""Triage a spreadsheet of appointment requests and produce a colour-coded Excel report.

Usage:
    python batch_triage.py sample_data/appointment_requests.csv
    python batch_triage.py requests.xlsx -o monday_report.xlsx --ai
"""
import argparse
import csv
import sys
import time
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from clinicpilot.ai import GeminiClient, triage_with_ai
from clinicpilot.queue_store import priority_key
from clinicpilot.rules import APPOINTMENTS, DURATIONS, LEVELS, REQUEST_TYPES, Intake
from clinicpilot.settings import load_env_file

LEVEL_COLOURS = {
    "emergency": ("C0392B", "FFFFFF"),
    "urgent": ("E67E22", "FFFFFF"),
    "routine": ("2E86C1", "FFFFFF"),
    "telehealth": ("7D3C98", "FFFFFF"),
    "selfcare": ("229954", "FFFFFF"),
}
HEADER_FILL = PatternFill("solid", fgColor="1E6A5B")
HEADER_FONT = Font(bold=True, color="FFFFFF")

REQUEST_ALIASES = {label.lower(): key for key, label in REQUEST_TYPES.items()}
REQUEST_ALIASES.update({
    "new problem": "new", "follow up": "followup", "follow-up": "followup", "prescription": "script",
    "repeat prescription": "script", "results": "results", "certificate": "certificate",
    "medical certificate": "certificate", "vaccination": "vaccine", "vaccine": "vaccine",
})
DURATION_ALIASES = {
    "<24h": "lt24", "less than 24 hours": "lt24", "today": "lt24",
    "1-3 days": "1to3", "1 to 3 days": "1to3",
    "4-7 days": "4to7", "4 to 7 days": "4to7",
    ">1 week": "gt7", "more than a week": "gt7", "over a week": "gt7",
    "n/a": "na", "": "na",
}


def read_rows(path):
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        headers = [str(h or "").strip().lower() for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(v not in (None, "") for v in row)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [{k.strip().lower(): v for k, v in row.items()} for row in csv.DictReader(f)]


def to_bool(value):
    return str(value or "").strip().lower() in ("yes", "y", "true", "1")


def normalise(row):
    request_type = str(row.get("request_type") or "new").strip().lower()
    duration = str(row.get("duration") or "").strip().lower()
    return {
        "name": str(row.get("name") or "").strip(),
        "age": row.get("age"),
        "request_type": REQUEST_ALIASES.get(request_type, request_type),
        "symptoms": str(row.get("symptoms") or ""),
        "duration": DURATION_ALIASES.get(duration, duration),
        "severity": row.get("severity") if row.get("severity") not in (None, "") else 3,
        "pregnant": to_bool(row.get("pregnant")),
        "chronic": to_bool(row.get("chronic")),
    }


def process(rows, client=None):
    done, problems = [], []
    for number, row in enumerate(rows, start=2):
        try:
            intake = Intake.from_dict(normalise(row))
        except ValueError as exc:
            problems.append({"row": number, "name": row.get("name") or "", "problem": str(exc)})
            continue
        result = triage_with_ai(intake, client)
        done.append({"intake": intake, "result": result, "row": number})
    done.sort(key=lambda d: priority_key({"level": d["result"].level, "severity": d["intake"].severity,
                                          "received": d["row"]}))
    return done, problems


def autosize(sheet, widths):
    for i, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width


def style_header(sheet):
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def write_report(done, problems, output, ai_used):
    wb = Workbook()

    queue = wb.active
    queue.title = "Triage queue"
    queue.append(["#", "Priority", "Patient", "Age", "Request", "Duration", "Severity",
                  "Booking", "Target wait", "Reception summary", "Reasons", "Decided by"])
    for i, item in enumerate(done, start=1):
        intake, result = item["intake"], item["result"]
        queue.append([
            i, LEVELS[result.level]["short"], intake.name, intake.age, REQUEST_TYPES[intake.request_type],
            DURATIONS[intake.duration], intake.severity, APPOINTMENTS[result.appointment],
            LEVELS[result.level]["wait"], result.summary,
            "; ".join(r["text"] for r in result.reasons),
            "Rules + Gemini" if result.source == "ai" else "Rules",
        ])
        fill, font = LEVEL_COLOURS[result.level]
        cell = queue.cell(row=i + 1, column=2)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(bold=True, color=font)
        for col in (10, 11):
            queue.cell(row=i + 1, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    style_header(queue)
    queue.auto_filter.ref = queue.dimensions
    autosize(queue, [5, 12, 22, 6, 16, 11, 9, 20, 30, 60, 60, 14])

    summary = wb.create_sheet("Summary")
    counts = Counter(item["result"].level for item in done)
    summary.append(["ClinicPilot triage report"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Generated", time.strftime("%Y-%m-%d %H:%M")])
    summary.append(["Mode", "Rules + Gemini AI" if ai_used else "Rules engine"])
    summary.append(["Requests triaged", len(done)])
    summary.append(["Rows needing review", len(problems)])
    summary.append([])
    summary.append(["Priority", "Requests", "Share"])
    header_row = summary.max_row
    for level, info in LEVELS.items():
        n = counts.get(level, 0)
        summary.append([info["short"], n, n / len(done) if done else 0])
        row = summary.max_row
        fill, font = LEVEL_COLOURS[level]
        summary.cell(row=row, column=1).fill = PatternFill("solid", fgColor=fill)
        summary.cell(row=row, column=1).font = Font(bold=True, color=font)
        summary.cell(row=row, column=3).number_format = "0%"
    for cell in summary[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    summary.append([])
    summary.append(["Booking type", "Requests"])
    booking_header = summary.max_row
    for key, label in APPOINTMENTS.items():
        summary.append([label, sum(1 for item in done if item["result"].appointment == key)])
    for cell in summary[booking_header]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    autosize(summary, [24, 22, 10])

    if problems:
        review = wb.create_sheet("Needs review")
        review.append(["Spreadsheet row", "Patient", "Problem"])
        for p in problems:
            review.append([p["row"], p["name"], p["problem"]])
        style_header(review)
        autosize(review, [16, 22, 60])

    wb.save(output)


def main(argv=None):
    parser = argparse.ArgumentParser(description="Triage a CSV or Excel file of appointment requests.")
    parser.add_argument("input", help="CSV or XLSX file with one request per row")
    parser.add_argument("-o", "--output", default="triage_report.xlsx", help="Excel file to create")
    parser.add_argument("--ai", action="store_true", help="also use Gemini (needs GEMINI_API_KEY)")
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f"File not found: {args.input}")
        return 1

    client = None
    if args.ai:
        load_env_file()
        client = GeminiClient()
        if not client.enabled:
            print("GEMINI_API_KEY is not set, so only the rules engine will be used.")

    started = time.time()
    rows = read_rows(args.input)
    done, problems = process(rows, client)
    write_report(done, problems, args.output, ai_used=bool(client and client.enabled))

    counts = Counter(item["result"].level for item in done)
    print(f"\nTriaged {len(done)} requests in {time.time() - started:.1f}s")
    for level, info in LEVELS.items():
        print(f"  {info['short']:<11} {counts.get(level, 0)}")
    if problems:
        print(f"  {len(problems)} row(s) need review (see the 'Needs review' sheet)")
    print(f"\nReport saved to {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
